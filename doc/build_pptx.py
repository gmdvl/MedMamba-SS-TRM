#!/usr/bin/env python3
"""Build a 16:9 PPTX from doc/PRESENTATION_MedMamba_to_TRM.md.

Layout is measured, not guessed: every block reports a height, the slide packs
them top-down, shrinks the type a step at a time if they don't fit, and spills
to a "(cont.)" slide only when shrinking is exhausted.
"""
import os, re, sys, glob
# python-pptx may live in a local vendor dir (PPTX_LIBS) or be installed normally.
for _cand in [os.environ.get('PPTX_LIBS'),
              os.path.join(os.path.dirname(os.path.abspath(__file__)), 'pylibs')]:
    if _cand and os.path.isdir(_cand):
        sys.path.insert(0, _cand); break

# Figure PNGs: rendered from doc/figures/*.svg (see README). Override with PPTX_PNG.


from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

DOC   = '/data/dante_data/documents/Masters/courses/Thesis/g-medmamba/doc/PRESENTATION_MedMamba_to_TRM.md'
PNG   = os.environ.get('PPTX_PNG',
        os.path.join(os.path.dirname(os.path.abspath(__file__)), 'figures', '_png'))
OUT   = '/data/dante_data/documents/Masters/courses/Thesis/g-medmamba/doc/PRESENTATION_MedMamba_to_TRM.pptx'

# palette from doc/figures/README.md
INK    = RGBColor(0x23, 0x28, 0x2f)
MUTED  = RGBColor(0x6b, 0x74, 0x82)
RULE   = RGBColor(0xd8, 0xdc, 0xe2)
GREYBL = RGBColor(0x5a, 0x64, 0x72)
ORANGE = RGBColor(0xc2, 0x70, 0x3a)
PURPLE = RGBColor(0x6a, 0x5a, 0x92)
GREEN  = RGBColor(0x4e, 0x8c, 0x6a)
RED    = RGBColor(0xb4, 0x55, 0x3f)
WHITE  = RGBColor(0xff, 0xff, 0xff)
BAND   = RGBColor(0xf4, 0xf5, 0xf7)

BODY_FONT = 'Calibri'
MONO_FONT = 'Consolas'

SW, SH   = 13.3333, 7.5           # inches
ML, MR   = 0.55, 0.55
TOP      = 1.22
BOT      = 7.02
CW       = SW - ML - MR           # content width

PART_COLOR = {1: ORANGE, 2: PURPLE, 3: GREYBL, 4: GREEN, 0: GREYBL}

# --------------------------------------------------------------------------- parse

def split_slides(md):
    lines = md.split('\n')
    chunks, cur = [], []
    for ln in lines:
        if ln.strip() == '---':
            chunks.append('\n'.join(cur)); cur = []
        else:
            cur.append(ln)
    chunks.append('\n'.join(cur))
    return [c for c in chunks if c.strip()]


def parse_blocks(body):
    """Turn slide markdown into a list of typed blocks."""
    blocks, i = [], 0
    lines = body.split('\n')
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1; continue
        # image
        m = re.match(r'!\[(.*?)\]\((.*?)\)', s)
        if m:
            blocks.append(('image', m.group(2))); i += 1; continue
        # code fence
        if s.startswith('```'):
            j = i + 1; code = []
            while j < len(lines) and not lines[j].strip().startswith('```'):
                code.append(lines[j]); j += 1
            blocks.append(('code', '\n'.join(code))); i = j + 1; continue
        # table
        if s.startswith('|'):
            j = i; rows = []
            while j < len(lines) and lines[j].strip().startswith('|'):
                rows.append(lines[j].strip()); j += 1
            cells = []
            for r in rows:
                if re.match(r'^\|[\s:\-|]+\|$', r):
                    continue
                parts = [c.strip() for c in r.strip('|').split('|')]
                cells.append(parts)
            if cells:
                blocks.append(('table', cells))
            i = j; continue
        # blockquote
        if s.startswith('>'):
            j = i; quote = []
            while j < len(lines) and lines[j].strip().startswith('>'):
                quote.append(lines[j].strip().lstrip('>').strip()); j += 1
            text = ' '.join(q for q in quote if q)
            if text.startswith('**Say:**'):
                blocks.append(('note', text[len('**Say:**'):].strip()))
            else:
                blocks.append(('quote', text))
            i = j; continue
        # list
        if re.match(r'^[-*]\s+', s) or re.match(r'^\d+\.\s+', s):
            j = i; items = []
            while j < len(lines):
                t = lines[j].strip()
                m2 = re.match(r'^(?:[-*]|\d+\.)\s+(.*)$', t)
                if m2:
                    items.append(m2.group(1)); j += 1
                elif t and not re.match(r'^(\||>|!\[|```|#)', t) and items:
                    items[-1] += ' ' + t; j += 1      # continuation line
                else:
                    break
            blocks.append(('list', items)); i = j; continue
        # paragraph
        j = i; para = []
        while j < len(lines):
            t = lines[j].strip()
            if not t or re.match(r'^(\||>|!\[|```|#)', t) or re.match(r'^(?:[-*]|\d+\.)\s+', t):
                break
            # a line opening with a status dingbat is its own line, not a
            # continuation - the source uses them as an inline yes/no pair
            if para and t[0] in MARKERS:
                break
            para.append(t); j += 1
        blocks.append(('para', ' '.join(para))); i = j
    return blocks


# Glyphs absent from Calibri/Carlito render as tofu boxes. U+2011 becomes a
# plain hyphen; the dingbats become a filled circle whose COLOUR carries the
# meaning, matching the figure palette (green = good, red = problem).
NBHYPHEN = '\u2011'
MARKERS = {'\u2705': GREEN, '\u274c': RED, '\u26a0': RED, '\u2605': ORANGE}
MARKER_CH = '\u25cf'

def sanitize(text):
    return (text.replace(NBHYPHEN, '-')
                .replace('\u23f1', '')
                .replace('\u2713', MARKER_CH)
                .replace('\u2717', MARKER_CH))


def split_markers(runs):
    """Give every dingbat its own run so it can be coloured independently."""
    out = []
    for t, b, i, m in runs:
        buf = ''
        for ch in t:
            if ch in MARKERS:
                if buf:
                    out.append((buf, b, i, m, None)); buf = ''
                out.append((MARKER_CH, True, False, False, MARKERS[ch]))
            else:
                buf += ch
        if buf:
            out.append((buf, b, i, m, None))
    return out


INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))')

def runs_of(text, bold=False, italic=False, mono=False):
    """Markdown inline -> [(text, bold, italic, mono)]. Recurses, so `code`
    nested inside **bold** keeps both attributes instead of leaking backticks."""
    out = []
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            out += runs_of(tok[2:-2], True, italic, mono)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            out.append((tok[1:-1], bold, italic, True))
        elif tok.startswith('[') and '](' in tok:
            out += runs_of(tok[1:tok.index('](')], bold, italic, mono)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            out += runs_of(tok[1:-1], bold, True, mono)
        else:
            out.append((tok, bold, italic, mono))
    return out


def plain(text):
    return ''.join(r[0] for r in runs_of(text))

# --------------------------------------------------------------------------- measure

def wrap_lines(text, fs, width_in, mono=False):
    cpl = max(8, int(width_in * 72 / (fs * (0.62 if mono else 0.525))))
    words, lines, cur = text.split(), 0, 0
    for w in words:
        add = len(w) + (1 if cur else 0)
        if cur + add > cpl and cur:
            lines += 1; cur = len(w)
        else:
            cur += add
    return max(1, lines + (1 if cur else 0))


def col_widths(cells, width_in):
    n = max(len(r) for r in cells)
    weight = []
    for c in range(n):
        longest = max((len(plain(r[c])) for r in cells if len(r) > c), default=1)
        weight.append(max(4, min(longest, 60)))
    tot = sum(weight)
    return [width_in * w / tot for w in weight]


def img_natural(val, width_in):
    p = os.path.join(PNG, os.path.basename(val).replace('.svg', '.png'))
    iw, ih = Image.open(p).size
    return width_in * ih / iw


def block_height(kind, val, fs, width_in, img_cap=3.30):
    """Height in inches."""
    if kind == 'image':
        return min(img_natural(val, width_in), img_cap) + 0.12
    if kind == 'para':
        return wrap_lines(plain(val), fs, width_in) * fs * 1.30 / 72 + 0.10
    if kind == 'quote':
        return wrap_lines(plain(val), fs - 0.5, width_in - 0.35) * fs * 1.30 / 72 + 0.24
    if kind == 'list':
        h = 0.0
        for it in val:
            h += wrap_lines(plain(it), fs, width_in - 0.30) * fs * 1.30 / 72 + 0.06
        return h + 0.08
    if kind == 'code':
        return len(sanitize(val).split('\n')) * (fs - 1) * 1.30 / 72 + 0.22
    if kind == 'table':
        tfs = fs - 2.0
        widths = col_widths(val, width_in)
        h = 0.0
        for r in val:
            rows = 1
            for c, cell in enumerate(r):
                rows = max(rows, wrap_lines(plain(cell), tfs, widths[c] - 0.14))
            h += max(0.25, rows * tfs * 1.28 / 72 + 0.13)
        return h + 0.10
    return 0.0

# --------------------------------------------------------------------------- render

def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def fill_para(p, text, fs, color=INK, bold_all=False):
    for t, b, i, m, mark in split_markers(runs_of(sanitize(text))):
        r = p.add_run(); r.text = t
        f = r.font
        f.size = Pt(fs); f.name = MONO_FONT if m else BODY_FONT
        f.bold = b or bold_all; f.italic = i
        f.color.rgb = mark if mark else (MUTED if (m and not b) else color)
        if m:
            f.size = Pt(fs - 0.5)


def draw_image(slide, val, x, y, w, cap=3.30):
    p = os.path.join(PNG, os.path.basename(val).replace('.svg', '.png'))
    iw, ih = Image.open(p).size
    h = w * ih / iw
    if h > cap:
        h = cap; w = h * iw / ih
    slide.shapes.add_picture(p, Inches(x + (CW - w) / 2), Inches(y), Inches(w), Inches(h))
    return h + 0.12


def draw_table(slide, cells, x, y, w, fs):
    tfs = fs - 2.0
    widths = col_widths(cells, w)
    ncol = len(widths)
    rows = [r + [''] * (ncol - len(r)) for r in cells]
    heights = []
    for r in rows:
        n = 1
        for c, cell in enumerate(r):
            n = max(n, wrap_lines(plain(cell), tfs, widths[c] - 0.14))
        heights.append(max(0.25, n * tfs * 1.28 / 72 + 0.13))
    total = sum(heights)
    shp = slide.shapes.add_table(len(rows), ncol, Inches(x), Inches(y), Inches(w), Inches(total))
    tbl = shp.table
    tbl.first_row = False
    for c, cw in enumerate(widths):
        tbl.columns[c].width = Inches(cw)
    for ri, r in enumerate(rows):
        tbl.rows[ri].height = Inches(heights[ri])
        for ci, cell in enumerate(r):
            tc = tbl.cell(ri, ci)
            tc.margin_left = tc.margin_right = Inches(0.07)
            tc.margin_top = tc.margin_bottom = Inches(0.03)
            tc.vertical_anchor = MSO_ANCHOR.MIDDLE
            tc.fill.solid()
            tc.fill.fore_color.rgb = BAND if ri == 0 else WHITE
            tf = tc.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            num = re.fullmatch(r'[\d.,%✅⚠❌\-—–\s]*', plain(cell))
            p.alignment = PP_ALIGN.RIGHT if (num and plain(cell).strip() and ci > 0) else PP_ALIGN.LEFT
            fill_para(p, cell, tfs, color=INK if ri else GREYBL, bold_all=(ri == 0))
    return total + 0.10


def draw_code(slide, code, x, y, w, fs):
    lines = sanitize(code).split('\n')
    h = len(lines) * (fs - 1) * 1.30 / 72 + 0.22
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # rectangle
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xf6, 0xf5, 0xfa)
    box.line.color.rgb = RGBColor(0xdd, 0xd8, 0xe8); box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.16); tf.margin_top = Inches(0.09)
    tf.margin_right = tf.margin_bottom = Inches(0.05)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT          # autoshape default is centred
        code_part, _, comment = ln.partition('#')
        for seg, col in ((code_part, INK), (('#' + comment) if comment else '', GREEN)):
            if not seg:
                continue
            r = p.add_run(); r.text = seg
            r.font.size = Pt(fs - 1); r.font.name = MONO_FONT; r.font.color.rgb = col
    return h + 0.14


def draw_quote(slide, text, x, y, w, fs):
    h = wrap_lines(plain(text), fs - 0.5, w - 0.35) * fs * 1.30 / 72 + 0.24
    bar = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(0.055), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = PURPLE
    bar.line.fill.background(); bar.shadow.inherit = False
    tf = textbox(slide, x + 0.20, y + 0.07, w - 0.25, h)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    fill_para(tf.paragraphs[0], text, fs - 0.5)
    return h + 0.10

# --------------------------------------------------------------------------- slides

def chrome(prs, title, time_str, part, star, page, part_name):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = PART_COLOR.get(part, GREYBL)
    if title:
        tx0 = ML + (0.40 if star else 0.0)
        tf = textbox(slide, tx0, 0.40, CW - (2.85 if len(time_str) > 10 else 1.6) - (tx0 - ML), 0.62)
        p = tf.paragraphs[0]
        if star:
            st = slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT,
                                        Inches(ML - 0.01), Inches(0.50), Inches(0.30), Inches(0.30))
            st.fill.solid(); st.fill.fore_color.rgb = accent
            st.line.fill.background(); st.shadow.inherit = False
        for t, b, i, m, mark in split_markers(runs_of(sanitize(title))):
            r = p.add_run(); r.text = t
            r.font.size = Pt(25 if m else 27); r.font.bold = True; r.font.italic = i
            r.font.name = MONO_FONT if m else BODY_FONT
            r.font.color.rgb = mark if mark else INK
        if time_str:
            wide = len(time_str) > 10          # "1:00 · cuttable" and friends
            cw_chip = 2.70 if wide else 1.45
            t = textbox(slide, SW - MR - cw_chip, 0.50, cw_chip, 0.46)
            tp = t.paragraphs[0]; tp.alignment = PP_ALIGN.RIGHT
            rr = tp.add_run(); rr.text = time_str
            rr.font.size = Pt(11 if wide else 13); rr.font.name = BODY_FONT
            rr.font.color.rgb = MUTED
        ln = slide.shapes.add_shape(1, Inches(ML), Inches(1.08), Inches(CW), Inches(0.028))
        ln.fill.solid(); ln.fill.fore_color.rgb = RULE
        ln.line.fill.background(); ln.shadow.inherit = False
        tab = slide.shapes.add_shape(1, Inches(ML), Inches(1.08), Inches(1.15), Inches(0.028))
        tab.fill.solid(); tab.fill.fore_color.rgb = accent
        tab.line.fill.background(); tab.shadow.inherit = False
    if page:
        f = textbox(slide, ML, SH - 0.42, CW, 0.26)
        p = f.paragraphs[0]
        r = p.add_run(); r.text = sanitize(f'{part_name}      {page}')
        r.font.size = Pt(10); r.font.name = BODY_FONT; r.font.color.rgb = RGBColor(0x9a, 0xa2, 0xad)
    return slide


def add_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def render_content(prs, title, time_str, star, blocks, part, page_ref, part_name):
    note = ' '.join(v for k, v in blocks if k == 'note')
    body = [(k, v) for k, v in blocks if k != 'note']

    # Choose the largest font scale that fits. Images are not given a fixed
    # slice: they take whatever vertical room the text blocks leave over, so a
    # figure slide with little text shows a big figure.
    avail = BOT - TOP
    n_img = sum(1 for k, _ in body if k == 'image')

    def fit(fs):
        text_h = sum(block_height(k, v, fs, CW) for k, v in body if k != 'image')
        if not n_img:
            return text_h, 3.30
        cap = max(1.60, min(4.85, (avail - text_h - 0.12 * n_img) / n_img))
        return text_h + sum(block_height('image', v, fs, CW, cap)
                            for k, v in body if k == 'image'), cap

    img_cap = 3.30
    for fs in (15.5, 14.5, 13.5, 12.5, 11.5, 10.5, 9.5):
        total, cap = fit(fs)
        if total <= avail:
            groups = [body]; img_cap = cap
            break
    else:
        fs = 10.5
        img_cap = 3.30
        groups, cur, h = [], [], 0.0
        for k, v in body:
            bh = block_height(k, v, fs, CW, img_cap)
            if cur and h + bh > BOT - TOP:
                groups.append(cur); cur, h = [], 0.0
            cur.append((k, v)); h += bh
        if cur:
            groups.append(cur)

    for gi, group in enumerate(groups):
        t = title if gi == 0 else f'{title} (cont.)'
        page_ref[0] += 1
        slide = chrome(prs, t, time_str if gi == 0 else '', part, star and gi == 0,
                       page_ref[0], part_name)
        y = TOP
        for k, v in group:
            if k == 'image':
                y += draw_image(slide, v, ML, y, CW, img_cap)
            elif k == 'table':
                y += draw_table(slide, v, ML, y, CW, fs)
            elif k == 'code':
                y += draw_code(slide, v, ML, y, CW, fs)
            elif k == 'quote':
                y += draw_quote(slide, v, ML, y, CW, fs)
            elif k == 'list':
                h = block_height(k, v, fs, CW)
                tf = textbox(slide, ML + 0.04, y, CW - 0.04, h)
                for i, it in enumerate(v):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.space_after = Pt(4)
                    b = p.add_run(); b.text = '•  '
                    b.font.size = Pt(fs); b.font.name = BODY_FONT
                    b.font.color.rgb = PART_COLOR.get(part, GREYBL); b.font.bold = True
                    fill_para(p, it, fs)
                y += h
            elif k == 'para':
                h = block_height(k, v, fs, CW)
                tf = textbox(slide, ML, y, CW, h)
                fill_para(tf.paragraphs[0], v, fs)
                y += h
        if gi == len(groups) - 1:
            add_notes(slide, note)
    return page_ref[0]


def main():
    md = open(DOC).read()
    chunks = split_slides(md)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)

    part, part_name, page = 0, 'Intro', [0]

    for idx, ch in enumerate(chunks):
        lines = [l for l in ch.split('\n')]
        first = next((l.strip() for l in lines if l.strip()), '')

        # ---- title slide
        if idx == 0:
            page[0] += 1
            slide = chrome(prs, '', '', 0, False, 0, '')
            band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(SW), Inches(SH))
            band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0xfa, 0xfb, 0xfc)
            band.line.fill.background(); band.shadow.inherit = False
            bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(SW), Inches(0.30))
            bar.fill.solid(); bar.fill.fore_color.rgb = PURPLE
            bar.line.fill.background(); bar.shadow.inherit = False
            tf = textbox(slide, 1.0, 2.15, SW - 2.0, 1.2)
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = 'From MedMamba to G-MedMamba-R'
            r.font.size = Pt(44); r.font.bold = True; r.font.name = BODY_FONT; r.font.color.rgb = INK
            tf2 = textbox(slide, 1.0, 3.35, SW - 2.0, 0.8)
            p2 = tf2.paragraphs[0]
            r2 = p2.add_run()
            r2.text = 'Two numbers we deleted from the weights'
            r2.font.size = Pt(21); r2.font.name = BODY_FONT; r2.font.color.rgb = GREYBL
            tf3 = textbox(slide, 1.0, 4.55, SW - 2.0, 1.6)
            for line in ['Audience: engineers who have already seen MedMamba',
                         'Length: 24:15 in the intended cut \u2014 timed from the speaker script, not estimated',
                         'Part 2 is a self\u2011contained TRM tutorial \u2014 ~13 minutes, and stands alone',
                         'You do not need to know TRM. Every term is defined on first use.']:
                p3 = tf3.add_paragraph()
                r3 = p3.add_run(); r3.text = sanitize(line)
                r3.font.size = Pt(14); r3.font.name = BODY_FONT; r3.font.color.rgb = MUTED
                p3.space_after = Pt(5)
            continue

        # ---- section divider
        if first.startswith('# '):
            m = re.match(r'#\s+Part\s+(\d)', first)
            part = int(m.group(1)) if m else part
            part_name = first.lstrip('# ').strip()
            page[0] += 1
            slide = chrome(prs, '', '', part, False, 0, '')
            accent = PART_COLOR.get(part, GREYBL)
            band = slide.shapes.add_shape(1, Inches(0), Inches(2.55), Inches(SW), Inches(0.10))
            band.fill.solid(); band.fill.fore_color.rgb = accent
            band.line.fill.background(); band.shadow.inherit = False
            tf = textbox(slide, 1.0, 2.90, SW - 2.0, 1.0)
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = sanitize(part_name)
            r.font.size = Pt(38); r.font.bold = True; r.font.name = BODY_FONT; r.font.color.rgb = INK
            rest = [l for l in lines if l.strip() and not l.strip().startswith('# ')]
            if rest:
                blocks = parse_blocks('\n'.join(rest))
                y = 4.05
                for k, v in blocks:
                    if k == 'note':
                        continue
                    txt = v if isinstance(v, str) else ''
                    if k in ('para', 'quote') and txt:
                        h = wrap_lines(plain(txt), 14, SW - 2.0) * 14 * 1.35 / 72 + 0.12
                        t = textbox(slide, 1.0, y, SW - 2.0, h)
                        fill_para(t.paragraphs[0], txt, 14, color=MUTED)
                        y += h + 0.10
            continue

        # ---- content slide
        heading = next((l for l in lines if l.strip().startswith('## ')), None)
        if heading is None:
            continue
        raw = heading.strip()[3:].strip()
        star = '★' in raw
        raw = raw.replace('★', '')
        tm = re.search(r'⏱\s*([\d:]+)', raw)
        time_str = tm.group(1) if tm else ''
        raw = re.sub(r'⏱\s*[\d:]+', '', raw)
        extra = ''
        em = re.search(r'\*\(([^)]*)\)\*', raw)
        if em:
            extra = em.group(1); raw = re.sub(r'\*\([^)]*\)\*', '', raw)
        title = re.sub(r'\s+', ' ', raw).strip(' \u2014-·')
        if extra:
            time_str = (time_str + '  ·  ' + extra).strip(' ·')

        body_lines = lines[lines.index(heading) + 1:]
        blocks = parse_blocks('\n'.join(body_lines))
        render_content(prs, title, time_str, star, blocks, part, page, part_name)

    prs.save(OUT)
    print(f'wrote {OUT}')
    print(f'slides: {len(prs.slides.__iter__.__self__._sldIdLst)}')

main()
